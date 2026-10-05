"""The built-in Word templates: used until an admin uploads one, and the starting point to edit.

`build_defaults()` (`manage.py report_templates`) writes `defaults/<report>.<language>.docx`,
headings in each language (right to left for Arabic), values as docxtpl tags (see
`context.VARIABLES`).
"""

from collections.abc import Callable, Sequence
from pathlib import Path

from django.conf import settings
from django.utils import translation
from django.utils.translation import gettext as _
from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..choices import ReportChoices
from .context import HEAD_COUNTS, MIX, TERMS

DEFAULTS_DIR = Path(__file__).resolve().parent / "defaults"

FONT = "Arial"  # has Arabic glyphs on Windows, macOS and LibreOffice


def default_path(report: str, language: str) -> Path:
    """The built-in template of `report` in `language`, else in the default language."""
    path = DEFAULTS_DIR / f"{report}.{language}.docx"

    if path.exists():
        return path

    return DEFAULTS_DIR / f"{report}.{settings.LANGUAGE_CODE}.docx"


def build_defaults(directory: Path = DEFAULTS_DIR) -> list[Path]:
    """Write every report's template in every language; return the files written."""
    directory.mkdir(parents=True, exist_ok=True)
    written = []

    for report in ReportChoices:
        for language, _name in settings.LANGUAGES:
            with translation.override(language):
                document = _BUILDERS[report](rtl=translation.get_language_bidi())

            path = directory / f"{report}.{language}.docx"
            document.save(str(path))
            written.append(path)

    return written


def _capacity(*, rtl: bool) -> DocumentType:
    document = _document(rtl=rtl, landscape=True)

    _paragraph(document, f"{_('capacity report')}: {{{{ chapter.name }}}}", rtl=rtl, style="Title")
    _paragraph(document, "{{ date }}", rtl=rtl)

    counts = [label for key, label in HEAD_COUNTS]
    keys = [key for key, _label in HEAD_COUNTS]

    _loop_table(
        document,
        headers=[
            "#",
            _("faculty"),
            *map(str, counts),
            _("PhD equivalents"),
            _("per PhD"),
            _("capacity"),
            _("current students"),
            _("max students"),
            _("free seats"),
            _("staff"),
            _("compliance"),
        ],
        cells=[
            "{{ f.index }}",
            "{{ f.name }}",
            *(f"{{{{ f.counted.{key} }}}}" for key in keys),
            "{{ f.phd_equivalents }}",
            "{{ f.students_per_phd }}",
            "{{ f.capacity }}",
            "{{ f.current_students }}",
            "{{ f.max_students }}",
            "{{ f.free_seats }}",
            "{{ f.staff_percentage }}",
            "{{ f.compliance }}",
        ],
        loop="f in faculties",
        rtl=rtl,
        font_size=8,
    )

    _paragraph(
        document,
        f"{_('chapter capacity')}: {{{{ chapter.capacity }}}} "
        f"({_('max students')}: {{{{ chapter.max_students_allowed }}}})",
        rtl=rtl,
        bold=True,
    )
    _paragraph(
        document,
        f"{_('current students')}: {{{{ chapter.current_students }}}} · "
        f"{_('free seats')}: {{{{ chapter.free_seats }}}}",
        rtl=rtl,
    )
    _paragraph(document, f"{_('compliance')}: {{{{ chapter.compliance }}}}", rtl=rtl)

    _paragraph(document, _("violations"), rtl=rtl, style="Heading 2")
    _paragraph(document, "{%p for f in faculties if f.violations %}", rtl=rtl)
    _paragraph(document, "{{ f.name }}", rtl=rtl, bold=True)
    _paragraph(document, "{%p for v in f.violations %}", rtl=rtl)
    _paragraph(document, "{{ v }}", rtl=rtl, style="List Bullet")
    _paragraph(document, "{%p endfor %}", rtl=rtl)
    _paragraph(document, "{%p endfor %}", rtl=rtl)

    return document


def _faculty_staff(*, rtl: bool) -> DocumentType:
    document = _document(rtl=rtl, landscape=False)

    _faculty_heading(document, rtl=rtl)

    _paragraph(
        document,
        f"{_('staff')} ({{{{ staff_count }}}})",
        rtl=rtl,
        style="Heading 2",
    )
    _loop_table(
        document,
        headers=[
            "#",
            _("name"),
            _("specialization"),
            _("specialization type"),
            _("degree"),
            _("terms"),
            _("status"),
        ],
        cells=[
            "{{ e.index }}",
            "{{ e.name }}",
            "{{ e.specialization }}",
            "{{ e.specialization_type }}",
            "{{ e.degree }}",
            "{{ e.terms }}",
            "{{ e.status }}",
        ],
        loop="e in staff",
        rtl=rtl,
        font_size=9,
    )

    _paragraph(document, _("head counts"), rtl=rtl, style="Heading 2")
    _loop_table(
        document,
        headers=["", _("signed"), _("counted")],
        cells=["{{ h.label }}", "{{ h.signed }}", "{{ h.counted }}"],
        loop="h in head_counts",
        rtl=rtl,
        font_size=9,
    )

    _faculty_violations(document, rtl=rtl)

    return document


def _capacity_pivot(*, rtl: bool) -> DocumentType:
    document = _document(rtl=rtl, landscape=False)

    _paragraph(document, f"{_('capacity pivot')}: {{{{ chapter.name }}}}", rtl=rtl, style="Title")
    _paragraph(document, "{{ date }}", rtl=rtl)

    _loop_table(
        document,
        headers=[
            "#",
            _("faculty"),
            _("PhD equivalents"),
            _("capacity"),
            _("current students"),
            _("free seats"),
            _("compliance"),
        ],
        cells=[
            "{{ f.index }}",
            "{{ f.name }}",
            "{{ f.phd_equivalents }}",
            "{{ f.capacity }}",
            "{{ f.current_students }}",
            "{{ f.free_seats }}",
            "{{ f.compliance }}",
        ],
        loop="f in faculties",
        rtl=rtl,
        font_size=9,
        footer=[
            "",
            _("chapter capacity"),
            "",
            "{{ chapter.capacity }}",
            "{{ chapter.current_students }}",
            "{{ chapter.free_seats }}",
            "{{ chapter.compliance }}",
        ],
    )

    # one small table per faculty: its PhDs by type (rows) and terms (columns)
    _paragraph(document, "{%p for f in faculties %}", rtl=rtl)
    _paragraph(document, "{{ f.name }}", rtl=rtl, style="Heading 2")
    _grid(
        document,
        [
            ["", *(str(label) for _key, label in TERMS), _("total")],
            *(
                [
                    str(label),
                    *(f"{{{{ f.{kind}.{key} }}}}" for key, _label in TERMS),
                    f"{{{{ f.{kind}.total }}}}",
                ]
                for kind, label in ((_SPECIALIZED, _("specialized")), (_SUPPORTED, _("supported")))
            ),
        ],
        rtl=rtl,
    )
    _paragraph(
        document,
        f"{_('masters')}: {{{{ f.masters }}}} · "
        f"{_('PhD equivalents')}: {{{{ f.phd_equivalents }}}} · "
        f"{_('students per PhD')}: {{{{ f.students_per_phd }}}} · "
        f"{_('staff')}: {{{{ f.staff_percentage }}}}",
        rtl=rtl,
    )
    _paragraph(
        document,
        f"{_('capacity')}: {{{{ f.capacity }}}} · "
        f"{_('current students')}: {{{{ f.current_students }}}} · "
        f"{_('max students')}: {{{{ f.max_students }}}} · "
        f"{_('free seats')}: {{{{ f.free_seats }}}} · "
        f"{{{{ f.compliance }}}}",
        rtl=rtl,
        bold=True,
    )
    _paragraph(document, _mix_line("f"), rtl=rtl)
    _paragraph(document, "{%p for v in f.violations %}", rtl=rtl)
    _paragraph(document, "{{ v }}", rtl=rtl, style="List Bullet")
    _paragraph(document, "{%p endfor %}", rtl=rtl)
    _paragraph(document, "{%p endfor %}", rtl=rtl)

    _paragraph(
        document, _("each cell: counted / signed, or one number when all are counted."), rtl=rtl
    )

    return document


def _audit(*, rtl: bool) -> DocumentType:
    document = _document(rtl=rtl, landscape=False)

    _paragraph(
        document, f"{_('calculation audit')}: {{{{ chapter.name }}}}", rtl=rtl, style="Title"
    )
    _paragraph(document, "{{ date }}", rtl=rtl)

    _paragraph(document, _("the counting rules"), rtl=rtl, style="Heading 2")
    _paragraph(document, "{%p for rule in rules %}", rtl=rtl)
    _paragraph(document, "{{ loop.index }}. {{ rule }}", rtl=rtl)
    _paragraph(document, "{%p endfor %}", rtl=rtl)

    steps = [_("step"), _("how"), _("result")]

    _paragraph(document, _("chapter capacity"), rtl=rtl, style="Heading 2")
    _loop_table(
        document,
        headers=["#", *steps],
        cells=["{{ s.index }}", "{{ s.label }}", "{{ s.how }}", "{{ s.value }}"],
        loop="s in steps",
        rtl=rtl,
        font_size=9,
    )

    _paragraph(document, "{%p for f in faculties %}", rtl=rtl)
    _paragraph(document, "{{ f.name }}", rtl=rtl, style="Heading 2")
    _loop_table(
        document,
        headers=["#", *steps],
        cells=["{{ s.index }}", "{{ s.label }}", "{{ s.how }}", "{{ s.value }}"],
        loop="s in f.steps",
        rtl=rtl,
        font_size=9,
    )
    _paragraph(document, _mix_line("f"), rtl=rtl)
    _paragraph(document, "{{ f.compliance }}", rtl=rtl, bold=True)
    _paragraph(document, "{%p for v in f.violations %}", rtl=rtl)
    _paragraph(document, "{{ v }}", rtl=rtl, style="List Bullet")
    _paragraph(document, "{%p endfor %}", rtl=rtl)
    _paragraph(document, "{%p if f.uncounted %}", rtl=rtl)
    _paragraph(document, _("contracts not counted"), rtl=rtl, bold=True)
    _loop_table(
        document,
        headers=["#", _("name"), _("specialization"), _("terms"), _("reason")],
        cells=[
            "{{ u.index }}",
            "{{ u.name }}",
            "{{ u.specialization }}",
            "{{ u.terms }}",
            "{{ u.reason }}",
        ],
        loop="u in f.uncounted",
        rtl=rtl,
        font_size=9,
    )
    _paragraph(document, "{%p endif %}", rtl=rtl)
    _paragraph(document, "{%p endfor %}", rtl=rtl)

    return document


def _faculty_staff_pivot(*, rtl: bool) -> DocumentType:
    document = _document(rtl=rtl, landscape=False)

    _faculty_heading(document, rtl=rtl)

    _paragraph(document, _("staff by specialization"), rtl=rtl, style="Heading 2")
    _loop_table(
        document,
        headers=[
            "#",
            _("specialization"),
            _("specialization type"),
            *(str(label) for _key, label in TERMS),
            _("masters"),
            _("total"),
        ],
        cells=[
            "{{ s.index }}",
            "{{ s.name }}",
            "{{ s.type }}",
            *(f"{{{{ s.{key} }}}}" for key, _label in TERMS),
            "{{ s.masters }}",
            "{{ s.total }}",
        ],
        loop="s in specializations",
        rtl=rtl,
        font_size=9,
        footer=[
            "",
            _("total"),
            "",
            *(f"{{{{ totals.{key} }}}}" for key, _label in TERMS),
            "{{ totals.masters }}",
            "{{ totals.total }}",
        ],
    )
    _paragraph(
        document, _("each cell: counted / signed, or one number when all are counted."), rtl=rtl
    )

    _faculty_violations(document, rtl=rtl)

    return document


_BUILDERS: dict[str, Callable[..., DocumentType]] = {
    ReportChoices.CAPACITY: _capacity,
    ReportChoices.FACULTY_STAFF: _faculty_staff,
    ReportChoices.CAPACITY_PIVOT: _capacity_pivot,
    ReportChoices.FACULTY_STAFF_PIVOT: _faculty_staff_pivot,
    ReportChoices.AUDIT: _audit,
}

_SPECIALIZED, _SUPPORTED = "specialized", "supported"


def _mix_line(owner: str) -> str:
    """`specialized 75% · supported 25% · ...`: the signed teachers as percentages."""
    shares = " · ".join(f"{label} {{{{ {owner}.mix.{key} }}}}" for key, label in MIX)

    return f"{_('signed teachers')}: {shares}"


def _faculty_heading(document: DocumentType, *, rtl: bool) -> None:
    """A faculty report's title and the faculty's numbers."""
    _paragraph(document, "{{ faculty.name }}", rtl=rtl, style="Title")
    _paragraph(
        document,
        f"{_('chapter')}: {{{{ chapter.name }}}} · {{{{ date }}}}",
        rtl=rtl,
    )

    _paragraph(document, _("capacity"), rtl=rtl, style="Heading 2")
    _key_values(
        document,
        [
            (_("capacity"), "{{ faculty.capacity }}"),
            (_("teaching capacity"), "{{ faculty.teaching_capacity }}"),
            (
                _("staff"),
                (
                    f"{{{{ faculty.staff_percentage }}}} "
                    f"({_('min')} {{{{ faculty.min_staff_percentage }}}})"
                ),
            ),
            *((str(label), f"{{{{ faculty.mix.{key} }}}}") for key, label in MIX),
            (_("students per PhD"), "{{ faculty.students_per_phd }}"),
            (_("current students"), "{{ faculty.current_students }}"),
            (_("free seats"), "{{ faculty.free_seats }}"),
            (_("compliance"), "{{ faculty.compliance }}"),
        ],
        rtl=rtl,
    )


def _faculty_violations(document: DocumentType, *, rtl: bool) -> None:
    _paragraph(document, "{%p if faculty.violations %}", rtl=rtl)
    _paragraph(document, _("violations"), rtl=rtl, style="Heading 2")
    _paragraph(document, "{%p for v in faculty.violations %}", rtl=rtl)
    _paragraph(document, "{{ v }}", rtl=rtl, style="List Bullet")
    _paragraph(document, "{%p endfor %}", rtl=rtl)
    _paragraph(document, "{%p endif %}", rtl=rtl)


def _document(*, rtl: bool, landscape: bool) -> DocumentType:
    document = Document()

    style = document.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:cs"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)

    section = document.sections[0]
    section.left_margin = section.right_margin = Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(1.5)

    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width

    if rtl:
        section._sectPr.append(OxmlElement("w:bidi"))  # noqa: SLF001

    return document


def _paragraph(
    document: DocumentType,
    text: str,
    *,
    rtl: bool,
    style: str | None = None,
    bold: bool = False,
) -> Paragraph:
    paragraph = document.add_paragraph(style=style)
    _text(paragraph, text, rtl=rtl, bold=bold)
    return paragraph


def _text(paragraph: Paragraph, text: str, *, rtl: bool, bold: bool = False, size: int = 0) -> None:
    run = paragraph.add_run(text)
    run.bold = bold or None

    if size:
        run.font.size = Pt(size)

    if rtl:
        paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))  # noqa: SLF001
        run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))  # noqa: SLF001


def _loop_table(  # noqa: PLR0913
    document: DocumentType,
    *,
    headers: Sequence[str],
    cells: Sequence[str],
    loop: str,
    rtl: bool,
    font_size: int,
    footer: Sequence[str] | None = None,
) -> Table:
    """A header row, then one row per item: docxtpl repeats the row between the `{%tr %}` rows.

    `footer`: a last, bold row (the totals).
    """
    table = document.add_table(rows=5 if footer else 4, cols=len(headers))
    table.style = "Table Grid"

    if rtl:
        table._tbl.tblPr.append(OxmlElement("w:bidiVisual"))  # noqa: SLF001

    _row(table, 0, headers, rtl=rtl, size=font_size, bold=True)
    _row(table, 1, [f"{{%tr for {loop} %}}"], rtl=rtl, size=font_size)
    _row(table, 2, cells, rtl=rtl, size=font_size)
    _row(table, 3, ["{%tr endfor %}"], rtl=rtl, size=font_size)

    if footer:
        _row(table, 4, footer, rtl=rtl, size=font_size, bold=True)

    return table


def _grid(document: DocumentType, rows: Sequence[Sequence[str]], *, rtl: bool) -> Table:
    """A small fixed table: its first row and first column are the (bold) headings."""
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"

    if rtl:
        table._tbl.tblPr.append(OxmlElement("w:bidiVisual"))  # noqa: SLF001

    for index, texts in enumerate(rows):
        _row(table, index, texts, rtl=rtl, size=9, bold=not index, bold_first=True)

    return table


def _key_values(document: DocumentType, rows: Sequence[tuple[str, str]], *, rtl: bool) -> Table:
    table = document.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"

    if rtl:
        table._tbl.tblPr.append(OxmlElement("w:bidiVisual"))  # noqa: SLF001

    for index, (label, value) in enumerate(rows):
        _row(table, index, [label, value], rtl=rtl, size=10, bold_first=True)

    return table


def _row(  # noqa: PLR0913
    table: Table,
    index: int,
    texts: Sequence[str],
    *,
    rtl: bool,
    size: int,
    bold: bool = False,
    bold_first: bool = False,
) -> None:
    for column, text in enumerate(texts):
        cell = table.rows[index].cells[column]
        _text(
            cell.paragraphs[0], text, rtl=rtl, size=size, bold=bold or (bold_first and not column)
        )
