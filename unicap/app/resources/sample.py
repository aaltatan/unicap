"""The sample file of an import: a few dummy rows to replace, with dropdowns in Excel.

```python
content = sample_file(EmployeeResource(chapter=chapter), "xlsx")
```

A column filled from a list (a choice, yes / no, another row's name) gets a dropdown of its
values; they are kept on a hidden sheet, so long lists (every employee) fit.
"""

import io

from django.utils import translation
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .base import TranslatedResource

BOM = "﻿"  # Excel then reads the CSV as UTF-8 (Arabic)

DROPDOWN_ROWS = 1000  # the rows under the header that get the dropdowns
LISTS_SHEET = "lists"
MIN_WIDTH = 14


def sample_file(resource: TranslatedResource, extension: str) -> bytes:
    """The resource's sample file, as `xlsx` (with dropdowns) or `csv`.

    Raises:
        ValueError: another extension.
    """
    match extension:
        case "csv":
            return (BOM + resource.sample_dataset().export("csv")).encode()
        case "xlsx":
            return _workbook(resource)

    msg = f"a sample is an xlsx or csv file, not {extension!r}"
    raise ValueError(msg)


def _workbook(resource: TranslatedResource) -> bytes:
    dataset = resource.sample_dataset()

    book = Workbook()

    sheet: Worksheet = book.active
    sheet.title = "sample"
    sheet.sheet_view.rightToLeft = translation.get_language_bidi()
    sheet.freeze_panes = "A2"

    sheet.append(dataset.headers)

    for row in dataset:
        sheet.append(row)

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        width = max(MIN_WIDTH, len(str(cell.value)) + 4)
        sheet.column_dimensions[cell.column_letter].width = width

    _add_dropdowns(book, sheet, resource)

    content = io.BytesIO()

    book.save(content)

    return content.getvalue()


def _add_dropdowns(book: Workbook, sheet: Worksheet, resource: TranslatedResource) -> None:
    """A dropdown on each listed column, its values on a hidden sheet."""
    options = resource.sample_options()

    if not options:
        return

    lists: Worksheet = book.create_sheet(LISTS_SHEET)
    lists.sheet_state = "hidden"

    listed = (
        (number, options[column])
        for number, column in enumerate(resource.sample_columns(), start=1)
        if column in options
    )

    for list_number, (number, values) in enumerate(listed, start=1):
        letter = get_column_letter(list_number)

        for row, value in enumerate(values, start=1):
            lists.cell(row=row, column=list_number, value=value)

        validation = DataValidation(
            type="list",
            formula1=f"={quote_sheetname(LISTS_SHEET)}!${letter}$1:${letter}${len(values)}",
            allow_blank=True,
            showErrorMessage=True,
        )

        target = get_column_letter(number)
        validation.add(f"{target}2:{target}{DROPDOWN_ROWS + 1}")

        sheet.add_data_validation(validation)
