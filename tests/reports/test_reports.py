"""Word reports: the downloads, the uploaded vs built-in templates, and the print page."""

from io import BytesIO
from pathlib import Path

import pytest
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse
from docx import Document
from pytest_django.fixtures import Settings
from selectolax.parser import HTMLParser

from unicap.app.choices import ReportChoices
from unicap.app.models import Chapter, Faculty, ReportTemplate
from unicap.app.reports import build_defaults
from unicap.app.reports.docx import DOCX_CONTENT_TYPE


@pytest.fixture(autouse=True)
def media(settings: Settings, tmp_path: Path) -> None:
    settings.MEDIA_ROOT = tmp_path


@pytest.fixture
def faculty(chapter: Chapter) -> Faculty:
    """A faculty of the sample chapter with contracts signed to it."""
    return Faculty.objects.for_chapter(chapter.pk).filter(contracts__isnull=False).first()


def docx_text(content: bytes) -> str:
    """Every paragraph and table cell of a .docx, one per line."""
    document = Document(BytesIO(content))
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([*(p.text for p in document.paragraphs), *cells])


def docx_upload(*paragraphs: str) -> SimpleUploadedFile:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return SimpleUploadedFile("template.docx", output.getvalue())


def faculty_names(chapter: Chapter) -> list[str]:
    return list(Faculty.objects.for_chapter(chapter.pk).values_list("name", flat=True))


def staff_names(faculty: Faculty) -> list[str]:
    return list(faculty.contracts.values_list("employee__name", flat=True))


# --- capacity report ----------------------------------------------------------------------


def test_capacity_report_downloads_a_docx(admin_client: Client, chapter: Chapter) -> None:
    response = admin_client.get(reverse("reports:capacity-docx"))

    assert response.status_code == 200
    assert response["Content-Type"] == DOCX_CONTENT_TYPE
    assert response["Content-Disposition"].startswith("attachment;")

    text = docx_text(b"".join(response.streaming_content))
    assert chapter.name in text
    assert all(name in text for name in faculty_names(chapter))


def test_capacity_report_compares_students_with_the_capacity(
    admin_client: Client, chapter: Chapter
) -> None:
    content = b"".join(admin_client.get(reverse("reports:capacity-docx")).streaming_content)

    text = docx_text(content)
    document = Document(BytesIO(content))
    headers = [cell.text for cell in document.tables[0].rows[0].cells]
    computer_science = [cell.text for cell in document.tables[0].rows[1].cells]

    start = headers.index("capacity")
    assert headers[start : start + 4] == [
        "capacity",
        "current students",
        "max students",
        "free seats",
    ]
    assert computer_science[start : start + 4] == ["75", "60", "200", "15"]
    assert "current students: 120 · free seats: 35" in text


def test_capacity_report_uses_the_language_of_the_request(
    admin_client: Client,
    chapter: Chapter,
) -> None:
    response = admin_client.get(reverse("reports:capacity-docx"), HTTP_ACCEPT_LANGUAGE="ar")

    text = docx_text(b"".join(response.streaming_content))
    assert "تقرير" in text  # the Arabic built-in template's title


def test_an_uploaded_template_replaces_the_built_in_one(
    admin_client: Client,
    chapter: Chapter,
) -> None:
    ReportTemplate.objects.create(
        report=ReportChoices.CAPACITY,
        language="en",
        file=docx_upload("custom: {{ chapter.name }} / {{ chapter.capacity }}"),
    )

    response = admin_client.get(reverse("reports:capacity-docx"))

    text = docx_text(b"".join(response.streaming_content))
    assert f"custom: {chapter.name} / " in text


def test_a_template_failing_to_render_shows_why(admin_client: Client, chapter: Chapter) -> None:
    ReportTemplate.objects.create(
        report=ReportChoices.CAPACITY,
        language="en",
        file=docx_upload("{{ chapter.name.missing() }}"),  # parses, but fails when filled
    )

    response = admin_client.get(reverse("reports:capacity-docx"), follow=True)

    assert response.redirect_chain == [(reverse("reports:capacity"), 302)]
    assert "Word template cannot be used" in response.content.decode()


def test_the_capacity_report_needs_view_permission(client: Client, chapter: Chapter) -> None:
    response = client.get(reverse("reports:capacity-docx"))

    assert response.status_code == 302
    assert response.url.startswith(reverse("login"))


# --- faculty staff ------------------------------------------------------------------------


def test_the_faculty_print_page_lists_its_staff(
    viewer_client: Client,
    faculty: Faculty,
) -> None:
    response = viewer_client.get(reverse("reports:staff", args=(faculty.pk,)))

    assert response.status_code == 200

    html = HTMLParser(response.content)
    rows = [row.text() for row in html.css("main tbody tr")]
    assert all(any(name in row for row in rows) for name in staff_names(faculty))
    assert html.css_first(f"a[href='{reverse('reports:staff-docx', args=(faculty.pk,))}']")


def test_the_faculty_staff_downloads_a_docx(admin_client: Client, faculty: Faculty) -> None:
    response = admin_client.get(reverse("reports:staff-docx", args=(faculty.pk,)))

    assert response.status_code == 200

    text = docx_text(b"".join(response.streaming_content))
    assert faculty.name in text
    assert all(name in text for name in staff_names(faculty))


def test_the_faculty_details_link_to_the_print_page(
    admin_client: Client,
    faculty: Faculty,
) -> None:
    response = admin_client.get(reverse("edu:faculties:details", args=(faculty.pk,)))

    html = HTMLParser(response.content)
    assert html.css_first(f"a[href='{reverse('reports:staff', args=(faculty.pk,))}']")


# --- templates ----------------------------------------------------------------------------


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("file", "reason"),
    [
        (SimpleUploadedFile("template.docx", b"not a zip"), "invalid_template"),
        (docx_upload("{{ chapter.name "), "invalid_template"),
        (SimpleUploadedFile("template.txt", b"text"), "invalid_extension"),
    ],
)
def test_an_invalid_template_is_refused(file: SimpleUploadedFile, reason: str) -> None:
    template = ReportTemplate(report=ReportChoices.CAPACITY, language="en", file=file)

    with pytest.raises(ValidationError) as error:
        template.full_clean()

    assert reason in [e.code for e in error.value.error_dict["file"]]


@pytest.mark.parametrize("report", ReportChoices.values)
def test_admins_download_the_built_in_templates(admin_client: Client, report: str) -> None:
    response = admin_client.get(reverse("reports:default-template", args=(report, "ar")))

    assert response.status_code == 200
    assert response["Content-Type"] == DOCX_CONTENT_TYPE


def test_built_in_templates_need_the_template_permission(viewer_client: Client) -> None:
    url = reverse("reports:default-template", args=(ReportChoices.CAPACITY, "en"))

    assert viewer_client.get(url).status_code == 403


def test_an_unknown_built_in_template_is_not_found(admin_client: Client) -> None:
    url = reverse("reports:default-template", args=("payroll", "en"))

    assert admin_client.get(url).status_code == 404


def test_build_defaults_writes_every_report_in_every_language(tmp_path: Path) -> None:
    written = build_defaults(tmp_path)

    assert sorted(path.name for path in written) == [
        "capacity.ar.docx",
        "capacity.en.docx",
        "capacity_audit.ar.docx",
        "capacity_audit.en.docx",
        "capacity_pivot.ar.docx",
        "capacity_pivot.en.docx",
        "faculty_staff.ar.docx",
        "faculty_staff.en.docx",
        "faculty_staff_pivot.ar.docx",
        "faculty_staff_pivot.en.docx",
    ]


def test_templates_cannot_reach_python_internals(admin_client: Client, chapter: Chapter) -> None:
    ReportTemplate.objects.create(
        report=ReportChoices.CAPACITY,
        language="en",
        file=docx_upload("{{ chapter.name.__class__.__mro__ }}"),  # sandboxed Jinja2
    )

    response = admin_client.get(reverse("reports:capacity-docx"), follow=True)

    assert response.redirect_chain == [(reverse("reports:capacity"), 302)]
    assert "Word template cannot be used" in response.content.decode()
