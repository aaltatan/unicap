"""The reports section: its sidebar and urls, each report's export menu, the audit and the mix."""

from io import BytesIO
from pathlib import Path

import pytest
from django.test import Client
from django.urls import reverse
from docx import Document
from pytest_django.fixtures import Settings
from selectolax.parser import HTMLParser

from unicap.app.choices import ReportChoices
from unicap.app.models import Chapter, Faculty
from unicap.app.reports import documents
from unicap.app.reports.docx import DOCX_CONTENT_TYPE
from unicap.app.templatetags.domain import mix_text
from unicap.domain import mix_of


@pytest.fixture(autouse=True)
def media(settings: Settings, tmp_path: Path) -> None:
    settings.MEDIA_ROOT = tmp_path


@pytest.fixture
def faculty(chapter: Chapter) -> Faculty:
    """A faculty of the sample chapter with contracts signed to it."""
    return Faculty.objects.for_chapter(chapter.pk).filter(contracts__isnull=False).first()


def docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([*(p.text for p in document.paragraphs), *cells])


def number(text: str) -> int:
    return int(text)


# --- the section --------------------------------------------------------------------------


def test_the_sidebar_has_a_reports_section(admin_client: Client, chapter: Chapter) -> None:
    tree = HTMLParser(admin_client.get(reverse("board:dashboard")).content)

    links = [a.attributes["href"] for a in tree.css("nav a.nav-link")]

    for name in ("reports:capacity", "reports:audit", "reports:faculties"):
        assert reverse(name) in links


def test_every_report_lives_under_reports(chapter: Chapter, faculty: Faculty) -> None:
    urls = [
        reverse("reports:capacity"),
        reverse("reports:capacity-pivot-pdf"),
        reverse("reports:audit-docx"),
        reverse("reports:faculties"),
        reverse("reports:staff", args=(faculty.pk,)),
        reverse("reports:staff-pivot-pdf", args=(faculty.pk,)),
    ]

    assert all(url.startswith("/reports/") for url in urls)
    assert reverse("reports:staff-pivot-pdf", args=(faculty.pk,)).endswith("/staff/pivot/pdf/")


def test_the_faculty_reports_page_lists_every_faculty(
    viewer_client: Client,
    chapter: Chapter,
) -> None:
    response = viewer_client.get(reverse("reports:faculties"))

    rows = HTMLParser(response.content).css("tbody tr")

    assert response.status_code == 200
    assert len(rows) == Faculty.objects.for_chapter(chapter.pk).count()


@pytest.mark.parametrize("name", ["reports:capacity", "reports:capacity-pivot", "reports:audit"])
def test_a_chapter_report_has_one_export_menu(
    admin_client: Client,
    chapter: Chapter,
    name: str,
) -> None:
    tree = HTMLParser(admin_client.get(reverse(name)).content)

    [menu] = tree.css("main [role='menu']")
    items = menu.css("[role='menuitem']")

    assert [item.attributes.get("href") for item in items] == [
        None,  # print: a button
        reverse(f"{name}-docx"),
        reverse(f"{name}-pdf"),
    ]
    assert "window.print()" in items[0].attributes["@click"]


@pytest.mark.parametrize(
    ("name", "other"),
    [("reports:staff", "reports:staff-pivot"), ("reports:staff-pivot", "reports:staff")],
)
def test_a_faculty_report_keeps_its_toggle_outside_the_menu(
    admin_client: Client,
    faculty: Faculty,
    name: str,
    other: str,
) -> None:
    tree = HTMLParser(admin_client.get(reverse(name, args=(faculty.pk,))).content)

    [menu] = tree.css("main [role='menu']")
    toggle = reverse(other, args=(faculty.pk,))

    assert menu.css_first(f"a[href='{reverse(f'{name}-pdf', args=(faculty.pk,))}']")
    assert menu.css_first(f"a[href='{toggle}']") is None
    assert tree.css_first(f"main a.btn[href='{toggle}']")


# --- the audit ----------------------------------------------------------------------------


def test_the_audit_explains_each_facultys_capacity(chapter: Chapter) -> None:
    report = Chapter.objects.get_snapshot(chapter.pk).report()

    context = documents.chapter_context(chapter, ReportChoices.AUDIT)

    assert len(context["rules"]) >= 8

    for row, faculty in zip(context["faculties"], report.faculties, strict=True):
        steps = {step["label"]: step for step in row["steps"]}
        counted = faculty.counted

        assert [step["index"] for step in row["steps"]] == list(range(1, len(row["steps"]) + 1))
        assert number(steps["contracts signed to the faculty"]["value"]) == len(faculty.statuses)
        assert number(steps["counted teachers"]["value"]) == len(faculty.counted_contracts)
        assert number(steps["counted PhDs"]["value"]) == counted.phds
        assert number(steps["PhD equivalents"]["value"]) == counted.phd_equivalents
        assert steps["PhD equivalents"]["how"] == f"{counted.phds} + {counted.masters_as_phds}"
        assert number(steps["teaching capacity"]["value"]) == (
            counted.phd_equivalents * faculty.faculty.students_per_phd
        )
        assert str(faculty.faculty.students_per_phd) in steps["teaching capacity"]["how"]
        assert number(steps["capacity"]["value"]) == faculty.capacity
        assert len(row["uncounted"]) == len(faculty.statuses) - len(faculty.counted_contracts)


def test_the_audit_adds_the_faculties_up_to_the_chapter(chapter: Chapter) -> None:
    report = Chapter.objects.get_snapshot(chapter.pk).report()

    [summed, capacity, *_rest] = documents.chapter_context(chapter, ReportChoices.AUDIT)["steps"]

    assert summed["how"] == " + ".join(str(f.capacity) for f in report.faculties)
    assert number(summed["value"]) == report.faculties_capacity
    assert number(capacity["value"]) == report.capacity


def test_the_audit_page_shows_the_rules_and_every_faculty(
    admin_client: Client,
    chapter: Chapter,
) -> None:
    response = admin_client.get(reverse("reports:audit"))

    tree = HTMLParser(response.content)
    faculties = Faculty.objects.for_chapter(chapter.pk).count()

    assert response.status_code == 200
    assert len(tree.css("main ol li")) >= 8
    assert len(tree.css("main table")) >= 1 + faculties  # the chapter's steps, then each one's


def test_the_audit_downloads_a_docx(admin_client: Client, chapter: Chapter) -> None:
    response = admin_client.get(reverse("reports:audit-docx"))

    assert response["Content-Type"] == DOCX_CONTENT_TYPE

    text = docx_text(b"".join(response.streaming_content))

    assert all(
        name in text
        for name in Faculty.objects.for_chapter(chapter.pk).values_list("name", flat=True)
    )
    assert "teaching capacity" in text
    assert "{{" not in text
    assert "{%" not in text


def test_the_audit_needs_view_permission(client: Client, chapter: Chapter) -> None:
    assert client.get(reverse("reports:audit")).status_code in (302, 403)


# --- the mix beside the staff percentage ----------------------------------------------------


def test_reports_carry_the_domains_mix(chapter: Chapter) -> None:
    report = Chapter.objects.get_snapshot(chapter.pk).report()

    context = documents.chapter_context(chapter, ReportChoices.CAPACITY)

    for row, faculty in zip(context["faculties"], report.faculties, strict=True):
        mix = mix_of(faculty.signed)

        assert faculty.mix == mix
        assert row["mix"]["specialized"] == f"{float(mix.specialized):.0f}%"
        assert row["mix"]["masters"] == f"{float(mix.masters):.0f}%"


def test_the_staff_percentage_tells_the_mix(admin_client: Client, chapter: Chapter) -> None:
    faculties = Faculty.objects.attach_reports(Faculty.objects.for_chapter(chapter.pk), chapter.pk)
    expected = {mix_text(faculty.report.mix) for faculty in faculties}

    table = HTMLParser(admin_client.get(reverse("edu:faculties:index")).content)
    report = HTMLParser(admin_client.get(reverse("reports:capacity")).content)

    assert {cell.attributes["title"] for cell in table.css('td[data-col="staff"]')} == expected
    assert expected <= {cell.attributes.get("title") for cell in report.css("main td")}


def test_the_faculty_modal_and_staff_report_draw_the_mix(
    admin_client: Client,
    faculty: Faculty,
) -> None:
    modal = admin_client.get(reverse("edu:faculties:details", args=(faculty.pk,)))
    page = admin_client.get(reverse("reports:staff", args=(faculty.pk,)))

    for response in (modal, page):
        text = HTMLParser(response.content).text()
        assert "Signed teachers" in text or "signed teachers" in text
        assert "parttime" in text


def test_the_api_returns_the_mix(chapter: Chapter) -> None:
    from unicap.app.serializers.capacity import _calculation_data  # noqa: PLC0415

    report = Chapter.objects.get_snapshot(chapter.pk).report().faculties[0]

    data = _calculation_data(report.calculation)

    assert data["mix"]["specialized"] + data["mix"]["supported"] in (0, 100)
    assert set(data["mix"]) == {
        "specialized",
        "supported",
        "fulltime",
        "parttime",
        "phds",
        "masters",
    }
