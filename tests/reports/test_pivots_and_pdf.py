"""The pivot reports (small specialized / supported tables) and every report as a PDF."""

import subprocess
from io import BytesIO
from pathlib import Path

import pytest
from django.test import Client
from django.urls import reverse
from docx import Document
from pytest_django.fixtures import Settings
from pytest_mock import MockerFixture
from selectolax.parser import HTMLParser

from unicap.app.choices import ReportChoices
from unicap.app.models import Chapter, Faculty
from unicap.app.reports import PdfError, documents, join, pdf, to_pdf
from unicap.app.reports.docx import DOCX_CONTENT_TYPE
from unicap.app.reports.pdf import PDF_CONTENT_TYPE

PDF = b"%PDF-1.7 converted"


@pytest.fixture(autouse=True)
def media(settings: Settings, tmp_path: Path) -> None:
    settings.MEDIA_ROOT = tmp_path


@pytest.fixture
def faculty(chapter: Chapter) -> Faculty:
    """A faculty of the sample chapter with contracts signed to it."""
    return Faculty.objects.for_chapter(chapter.pk).filter(contracts__isnull=False).first()


@pytest.fixture
def converter(mocker: MockerFixture) -> list[bytes]:
    """An office suite that writes a PDF; holds the .docx files it was given."""
    received: list[bytes] = []

    def convert(source: Path, target: Path) -> None:
        received.append(source.read_bytes())
        target.write_bytes(PDF)

    mocker.patch.object(pdf, "_converter", return_value=convert)

    return received


def docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([*(p.text for p in document.paragraphs), *cells])


def split(cell: str) -> tuple[int, int]:
    """A pivot cell as (counted, signed): `3` or `3 / 5`."""
    counted, _slash, signed = cell.partition(" / ")
    return int(counted), int(signed or counted)


# --- pdf ----------------------------------------------------------------------------------


@pytest.mark.parametrize(
    "url_name",
    ["reports:capacity-pdf", "reports:capacity-pivot-pdf"],
)
def test_a_chapter_report_downloads_as_pdf(
    admin_client: Client,
    chapter: Chapter,
    converter: list[bytes],
    url_name: str,
) -> None:
    response = admin_client.get(reverse(url_name))

    assert response.status_code == 200
    assert response["Content-Type"] == PDF_CONTENT_TYPE
    assert ".pdf" in response["Content-Disposition"]
    assert b"".join(response.streaming_content) == PDF
    assert chapter.name in docx_text(converter[0])  # the filled Word report was converted


@pytest.mark.parametrize("url_name", ["reports:staff-pdf", "reports:staff-pivot-pdf"])
def test_a_faculty_report_downloads_as_pdf(
    admin_client: Client,
    faculty: Faculty,
    converter: list[bytes],
    url_name: str,
) -> None:
    response = admin_client.get(reverse(url_name, args=(faculty.pk,)))

    assert response["Content-Type"] == PDF_CONTENT_TYPE
    assert b"".join(response.streaming_content) == PDF
    assert faculty.name in docx_text(converter[0])


def test_the_same_report_is_converted_once(
    admin_client: Client,
    chapter: Chapter,
    converter: list[bytes],
) -> None:
    for _again in range(3):
        response = admin_client.get(reverse("reports:capacity-pdf"))
        assert b"".join(response.streaming_content) == PDF

    assert len(converter) == 1


def test_a_change_in_the_chapter_is_converted_again(
    admin_client: Client,
    chapter: Chapter,
    converter: list[bytes],
) -> None:
    admin_client.get(reverse("reports:capacity-pdf"))

    faculty = Faculty.objects.for_chapter(chapter.pk).first()
    faculty.students_per_phd += 1
    faculty.save()

    admin_client.get(reverse("reports:capacity-pdf"))
    admin_client.get(reverse("reports:capacity-pivot-pdf"))  # another report: its own PDF

    assert len(converter) == 3


def test_a_failed_conversion_is_not_remembered(
    admin_client: Client,
    chapter: Chapter,
    mocker: MockerFixture,
) -> None:
    def fail(source: Path, target: Path) -> None:
        raise subprocess.CalledProcessError(1, "soffice")

    mocker.patch.object(pdf, "_converter", return_value=fail)
    admin_client.get(reverse("reports:capacity-pdf"))

    def convert(source: Path, target: Path) -> None:
        target.write_bytes(PDF)

    mocker.patch.object(pdf, "_converter", return_value=convert)
    response = admin_client.get(reverse("reports:capacity-pdf"))

    assert b"".join(response.streaming_content) == PDF


def test_without_an_office_suite_the_page_says_why(
    admin_client: Client,
    chapter: Chapter,
    mocker: MockerFixture,
) -> None:
    mocker.patch.object(pdf, "soffice_path", return_value=None)
    mocker.patch.object(pdf.shutil, "which", return_value=None)

    response = admin_client.get(reverse("reports:capacity-pdf"), follow=True)

    assert response.redirect_chain[0][0] == reverse("reports:capacity")
    assert "LibreOffice" in response.content.decode()


def test_a_failed_conversion_is_a_pdf_error(mocker: MockerFixture) -> None:
    def fail(source: Path, target: Path) -> None:
        raise subprocess.CalledProcessError(1, "soffice")

    mocker.patch.object(pdf, "_converter", return_value=fail)

    with pytest.raises(PdfError, match="cannot be made"):
        to_pdf(b"docx")


def test_libreoffice_is_preferred_and_configurable(
    settings: Settings,
    mocker: MockerFixture,
) -> None:
    settings.SOFFICE_PATH = "/opt/office/soffice"
    run = mocker.patch.object(pdf.subprocess, "run")

    with pytest.raises(PdfError):  # the mocked run writes no file
        to_pdf(b"docx")

    command = run.call_args.args[0]
    assert command[0] == "/opt/office/soffice"
    assert "--convert-to" in command


def test_a_pdf_needs_view_permission(client: Client, chapter: Chapter) -> None:
    assert client.get(reverse("reports:capacity-pdf")).status_code in (302, 403)


# --- capacity pivot -----------------------------------------------------------------------


def test_the_capacity_pivot_page_has_a_table_per_faculty(
    admin_client: Client,
    chapter: Chapter,
) -> None:
    response = admin_client.get(reverse("reports:capacity-pivot"))

    tree = HTMLParser(response.content)
    names = [heading.text(strip=True) for heading in tree.css("section h3")]

    assert names == [
        f.faculty.name for f in Chapter.objects.get_snapshot(chapter.pk).report().faculties
    ]
    assert len(tree.css("section table")) == len(names)
    assert reverse("reports:capacity-pivot-pdf") in response.content.decode()


def test_the_capacity_report_links_to_its_pivot_and_pdf(
    admin_client: Client,
    chapter: Chapter,
) -> None:
    content = admin_client.get(reverse("reports:capacity")).content.decode()

    assert reverse("reports:capacity-pivot") in content
    assert reverse("reports:capacity-pdf") in content


def test_the_capacity_pivot_holds_the_reports_head_counts(chapter: Chapter) -> None:
    report = Chapter.objects.get_snapshot(chapter.pk).report()

    context = documents.chapter_context(chapter, ReportChoices.CAPACITY_PIVOT)

    for row, faculty in zip(context["faculties"], report.faculties, strict=True):
        counted, signed = faculty.counted, faculty.signed

        assert split(row["specialized"]["fulltime_staff"]) == (
            counted.specialized_fulltime_staff,
            signed.specialized_fulltime_staff,
        )
        assert split(row["supported"]["parttime"]) == (
            counted.supported_parttime,
            signed.supported_parttime,
        )
        assert split(row["masters"]) == (counted.masters, signed.masters)

        totals = [split(row[kind]["total"]) for kind in ("specialized", "supported")]
        assert sum(c for c, _s in totals) == counted.phds
        assert sum(s for _c, s in totals) == signed.phds


def test_the_capacity_pivot_downloads_a_docx(admin_client: Client, chapter: Chapter) -> None:
    response = admin_client.get(reverse("reports:capacity-pivot-docx"))

    assert response["Content-Type"] == DOCX_CONTENT_TYPE

    content = b"".join(response.streaming_content)
    text = docx_text(content)
    faculties = list(Faculty.objects.for_chapter(chapter.pk).values_list("name", flat=True))

    assert all(name in text for name in faculties)
    assert len(Document(BytesIO(content)).tables) == 1 + len(faculties)  # the summary, then each
    assert "{{" not in text


# --- faculty staff pivot ------------------------------------------------------------------


def test_the_faculty_pivot_counts_its_staff_by_specialization(
    chapter: Chapter,
    faculty: Faculty,
) -> None:
    [faculty] = Faculty.objects.attach_reports([faculty], chapter.pk)
    report = faculty.report

    context = documents.faculty_report_context(chapter, faculty, ReportChoices.FACULTY_STAFF_PIVOT)

    rows = context["specializations"]
    assert [row["name"] for row in rows] == [s.name for s in report.faculty.faculty.specializations]
    assert {row["type"] for row in rows} <= {"specialized", "supported"}

    counted, signed = split(context["totals"]["total"])
    assert counted == len(report.counted_contracts)
    assert signed == len(report.included)
    assert sum(split(row["total"])[1] for row in rows) == signed
    assert split(context["totals"]["masters"]) == (report.counted.masters, report.signed.masters)


def test_the_faculty_pivot_page_and_docx(admin_client: Client, faculty: Faculty) -> None:
    page = admin_client.get(reverse("reports:staff-pivot", args=(faculty.pk,)))
    file = admin_client.get(reverse("reports:staff-pivot-docx", args=(faculty.pk,)))

    tree = HTMLParser(page.content)
    names = [row.css("td")[1].text(strip=True) for row in tree.css("tbody tr")]
    text = docx_text(b"".join(file.streaming_content))

    assert page.status_code == 200
    assert names
    assert all(name in text for name in names)
    assert faculty.name in text
    assert "{{" not in text


def test_the_staff_report_names_each_teachers_type(admin_client: Client, faculty: Faculty) -> None:
    response = admin_client.get(reverse("reports:staff-docx", args=(faculty.pk,)))

    document = Document(BytesIO(b"".join(response.streaming_content)))
    staff = next(table for table in document.tables if len(table.columns) == 7)
    types = {row.cells[3].text for row in staff.rows[1:]}

    assert staff.rows[0].cells[3].text == "specialization type"
    assert types <= {"specialized", "supported", ""}
    assert types & {"specialized", "supported"}


# --- every faculty's report in one file -----------------------------------------------------


def faculty_names(chapter: Chapter) -> list[str]:
    return list(Faculty.objects.for_chapter(chapter.pk).values_list("name", flat=True))


def page_breaks(content: bytes) -> list[str]:
    """The text of each paragraph that starts a new page."""
    document = Document(BytesIO(content))

    return [p.text for p in document.paragraphs if p.paragraph_format.page_break_before]


@pytest.mark.parametrize("url_name", ["reports:all-staff-docx", "reports:all-staff-pivot-docx"])
def test_every_facultys_report_downloads_as_one_word_file(
    admin_client: Client, chapter: Chapter, url_name: str
) -> None:
    response = admin_client.get(reverse(url_name))

    content = b"".join(response.streaming_content)
    text = docx_text(content)
    names = faculty_names(chapter)

    assert response["Content-Type"] == DOCX_CONTENT_TYPE
    assert "all_faculties" in response["Content-Disposition"]
    assert ".docx" in response["Content-Disposition"]
    # every faculty, in the tables' order, each after the first on a page of its own
    assert [text.index(name) for name in names] == sorted(text.index(name) for name in names)
    assert len(page_breaks(content)) == len(names) - 1
    assert all(any(name in start for name in names[1:]) for start in page_breaks(content))


def test_the_one_file_holds_what_each_facultys_own_file_holds(
    admin_client: Client, chapter: Chapter
) -> None:
    together = docx_text(
        b"".join(admin_client.get(reverse("reports:all-staff-docx")).streaming_content)
    )

    for faculty in Faculty.objects.for_chapter(chapter.pk):
        own = admin_client.get(reverse("reports:staff-docx", args=(faculty.pk,)))

        for line in docx_text(b"".join(own.streaming_content)).splitlines():
            assert line in together

    # every teacher signed to a faculty is in it once
    for name in chapter.contracts.signed().values_list("employee__name", flat=True):
        assert together.count(f"\n{name}\n") == 1


@pytest.mark.parametrize("url_name", ["reports:all-staff-pdf", "reports:all-staff-pivot-pdf"])
def test_every_facultys_report_downloads_as_one_pdf(
    admin_client: Client, chapter: Chapter, converter: list[bytes], url_name: str
) -> None:
    for _again in range(2):
        response = admin_client.get(reverse(url_name))

        assert response["Content-Type"] == PDF_CONTENT_TYPE
        assert b"".join(response.streaming_content) == PDF

    # one conversion, of the one file holding every faculty (and it is remembered)
    assert len(converter) == 1
    assert all(name in docx_text(converter[0]) for name in faculty_names(chapter))


def test_a_chapter_without_faculties_has_nothing_to_export(
    admin_client: Client, chapter: Chapter
) -> None:
    chapter.contracts.all().delete()
    chapter.employees.all().delete()
    chapter.faculties.all().delete()

    page = HTMLParser(admin_client.get(reverse("reports:faculties")).content)
    response = admin_client.get(reverse("reports:all-staff-docx"), follow=True)

    assert page.css_first("[data-export-all]") is None
    assert response.redirect_chain[-1][0] == reverse("reports:faculties")
    assert "no faculty yet" in response.content.decode()


def test_the_faculties_page_exports_them_all_from_one_menu(
    admin_client: Client, viewer_client: Client, chapter: Chapter
) -> None:
    html = HTMLParser(admin_client.get(reverse("reports:faculties")).content)

    links = [a.attributes["href"] for a in html.css("[data-export-all] [role=menu] a")]

    assert links == [
        reverse(name)
        for name in (
            "reports:all-staff-docx",
            "reports:all-staff-pdf",
            "reports:all-staff-pivot-docx",
            "reports:all-staff-pivot-pdf",
        )
    ]
    assert viewer_client.get(links[0]).status_code == 200  # whoever may see the reports


def test_all_faculties_need_view_permission(client: Client, chapter: Chapter) -> None:
    response = client.get(reverse("reports:all-staff-docx"))

    assert response.status_code == 302
    assert "/login" in response["Location"] or "accounts" in response["Location"]


def test_join_keeps_the_first_documents_page_setup() -> None:
    def document(*paragraphs: str) -> bytes:
        made = Document()
        for text in paragraphs:
            made.add_paragraph(text)
        output = BytesIO()
        made.save(output)
        return output.getvalue()

    joined = Document(BytesIO(join([document("one", "1"), document("two"), document("three")])))

    assert [p.text for p in joined.paragraphs] == ["one", "1", "two", "three"]
    assert [p.text for p in joined.paragraphs if p.paragraph_format.page_break_before] == [
        "two",
        "three",
    ]
    assert len(joined.sections) == 1
    assert joined.element.body[-1].tag.endswith("sectPr")  # still the body's last element
    assert join([document("alone")]) is not None

    with pytest.raises(ValueError, match="at least one"):
        join([])


def test_a_joined_document_starting_with_a_table_still_breaks_the_page() -> None:
    def document(cell: str) -> bytes:
        made = Document()
        made.add_table(rows=1, cols=1).cell(0, 0).text = cell
        output = BytesIO()
        made.save(output)
        return output.getvalue()

    joined = Document(BytesIO(join([document("a"), document("b")])))

    assert [table.cell(0, 0).text for table in joined.tables] == ["a", "b"]
    assert len([p for p in joined.paragraphs if p.paragraph_format.page_break_before]) == 1
